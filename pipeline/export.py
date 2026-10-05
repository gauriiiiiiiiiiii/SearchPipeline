"""Stage 3: build events.xlsx / events.csv / search_results.csv from the DB."""

import csv
import datetime as dt
import json
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import normalize
from .settings import EVENTS_CSV, EVENTS_XLSX, KEEP_OUTSIDE_INDIA, LINKS_CSV
from .util import log

# explicit FF alpha: some viewers render 00-alpha colours as transparent
HEAD_FONT = Font(name="Arial", bold=True, color="FF000000")
HEAD_FILL = PatternFill("solid", fgColor="FFBDD7EE", bgColor="FFBDD7EE")
BODY_FONT = Font(name="Arial", size=10, color="FF000000")
TITLE_FONT = Font(name="Arial", bold=True, size=14, color="FF1F4E78")

WIDTHS = {"Status": 11, "Event Name": 48, "Start Date": 11, "Day": 10, "Start Time": 10,
          "End Date": 11, "End Time": 10, "Duration (hrs)": 9, "Event Format": 10,
          "Venue": 30, "Area / Locality": 28, "City": 12, "State": 14, "Country": 8,
          "Full Address": 44, "Latitude": 10, "Longitude": 10, "Location Check": 12,
          "Hosted By": 26, "Organizer Link": 36, "Event Type": 14,
          "Industry / Sector Focus": 16, "Keywords": 30, "Startup Relevant": 9,
          "Price": 14, "Registration Link": 40, "Event Link": 50, "Event Status": 11,
          "Description": 60, "Image": 30, "Source Platform": 20, "Found On": 50,
          "Search Queries": 44, "Search Cities": 14, "Extraction Method": 10,
          "Confidence": 10}

LINK_COLS = ["engine", "city", "template", "query", "page", "rank", "title", "url",
             "domain", "crawl_status", "events_found"]


def _safe(v):
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def _sheet(ws, header, rows, widths=None):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    for r in rows:
        ws.append([_safe(v) for v in r])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, 16)
    ws.freeze_panes = "B2"


def collect(db):
    """Every stored event as an output row, tagged with the search that found it."""
    queries, cities = defaultdict(set), defaultdict(set)
    for url, q, c in db.execute("SELECT url, query, city FROM results"):
        queries[url].add(q)
        cities[url].add(c)
    parent = dict(db.execute("SELECT url, parent FROM pages"))
    now = dt.datetime.now()

    rows = []
    for page_url, data in db.execute("SELECT page_url, data FROM events ORDER BY id"):
        raw = json.loads(data)
        root = page_url if page_url in queries else (parent.get(page_url) or page_url)
        raw["found_on"] = root
        raw["queries"] = " | ".join(sorted(queries.get(root, ())))
        raw["search_cities"] = " | ".join(sorted(cities.get(root, ())))
        rows.append(normalize.row(raw, now))
    return rows


def run(db):
    raw_rows = collect(db)
    events, stats = normalize.merge(raw_rows, KEEP_OUTSIDE_INDIA)
    cols = normalize.COLUMNS

    with open(EVENTS_CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(events)

    crawl = dict(((u, (s, n)) for u, s, n in db.execute(
        "SELECT url, status, n_events FROM pages")))
    found_via = Counter()
    for u, p, n in db.execute("SELECT url, parent, n_events FROM pages WHERE parent IS NOT NULL"):
        found_via[p] += n or 0
    links = []
    for r in db.execute("SELECT engine, city, template, query, page, rank, title, url, domain "
                        "FROM results ORDER BY city, template, engine, rank"):
        st, n = crawl.get(r[7], ("not visited", ""))
        links.append(list(r) + [st, (n or 0) + found_via[r[7]] if st != "not visited" else ""])
    with open(LINKS_CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(LINK_COLS)
        w.writerows(links)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    by_status = Counter(e["Status"] for e in events)
    lines = [
        ("Generated", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("", ""),
        ("Search queries run", db.execute("SELECT COUNT(DISTINCT query) FROM done").fetchone()[0]),
        ("Result links", db.execute("SELECT COUNT(DISTINCT url) FROM results").fetchone()[0]),
        ("Pages crawled (incl. followed links)", len(crawl)),
        ("Events extracted (raw)", len(raw_rows)),
        ("Dropped: noise titles", stats["noise"]),
        ("Dropped: outside India", stats["outside_india"]),
        ("Merged duplicates", stats["duplicates"]),
        ("Unique events", len(events)),
        ("", ""),
        ("Upcoming", by_status["Upcoming"] + by_status["Ongoing"]),
        ("Past", by_status["Past"]),
        ("Date unknown", by_status["Date unknown"]),
        ("Startup relevant", sum(e["Startup Relevant"] == "Yes" for e in events)),
        ("", ""),
        ("Events by city", ""),
    ] + [(f"   {c or '(not found)'}", n) for c, n in
         Counter(e["City"] for e in events).most_common()] + [
        ("", ""), ("Events by platform", ""),
    ] + [(f"   {p}", n) for p, n in Counter(
        e["Source Platform"].split(" | ")[0] for e in events).most_common(15)]
    ws["A1"] = "Event pipeline summary"
    ws["A1"].font = TITLE_FONT
    for i, (k, v) in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=k).font = (HEAD_FONT if k and not k.startswith(" ")
                                                  and v == "" else BODY_FONT)
        ws.cell(row=i, column=2, value=v).font = BODY_FONT
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 20

    up = [e for e in events if e["Status"] in ("Upcoming", "Ongoing")]
    past = [e for e in events if e["Status"] == "Past"]
    for title, subset in (("Upcoming events", up), ("Past events", past), ("All events", events)):
        _sheet(wb.create_sheet(title), cols, [[e[c] for c in cols] for e in subset], WIDTHS)
    _sheet(wb.create_sheet("Search links"), LINK_COLS, links,
           {"engine": 8, "city": 12, "template": 32, "query": 42, "page": 6, "rank": 6,
            "title": 46, "url": 60, "domain": 24, "crawl_status": 12, "events_found": 12})
    _sheet(wb.create_sheet("Crawl log"),
           ["url", "found_from", "status", "fetched_via", "http", "events", "note", "ts"],
           db.execute("SELECT url, parent, status, via, http, n_events, note, ts FROM pages "
                      "ORDER BY ts").fetchall(),
           {"url": 60, "found_from": 50, "status": 10, "fetched_via": 10, "http": 6,
            "events": 7, "note": 30, "ts": 20})
    wb.save(EVENTS_XLSX)
    log(f"[export] {len(events)} unique events ({len(up)} upcoming, {len(past)} past) -> "
        f"output/events.xlsx, events.csv, search_results.csv")
